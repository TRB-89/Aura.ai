"""
Word document creation service using python-docx.
Falls back to this when Microsoft Word COM is unavailable.
"""
import os
import logging
from pathlib import Path
from typing import Optional
from datetime import datetime

logger = logging.getLogger("aura.office.word")


def get_output_dir() -> Path:
    """Get the approved AURA output directory."""
    output = Path.home() / "Documents" / "AURA"
    output.mkdir(parents=True, exist_ok=True)
    return output


async def create_word_document(
    topic: str,
    doc_type: str = "assignment",
    filename: Optional[str] = None,
    sections: Optional[list] = None,
) -> dict:
    """Create a properly formatted Word document on the topic provided."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.style import WD_STYLE_TYPE
    except ImportError:
        return {"success": False, "error": "python-docx is not installed"}

    # ── Generate content ────────────────────────────────────────────────────
    content = _generate_document_content(topic, doc_type, sections)
    
    # ── Build filename ───────────────────────────────────────────────────────
    if not filename:
        safe_topic = "".join(c if c.isalnum() or c in " _-" else "" for c in topic)[:50]
        filename = f"{safe_topic.strip()}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    
    output_path = get_output_dir() / f"{filename}.docx"
    
    # ── Create document ──────────────────────────────────────────────────────
    doc = Document()
    
    # Page margins
    for section in doc.sections:
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
    
    # Title
    title_para = doc.add_heading(content["title"], level=0)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.runs[0] if title_para.runs else title_para.add_run(content["title"])
    title_run.font.size = Pt(20)
    title_run.bold = True
    
    # Subtitle / metadata
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"Type: {doc_type.title()} | Date: {datetime.now().strftime('%B %d, %Y')}")
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    
    doc.add_paragraph()  # spacing
    
    # Sections
    for section_data in content["sections"]:
        # Heading
        heading = doc.add_heading(section_data["heading"], level=1)
        heading_run = heading.runs[0] if heading.runs else heading.add_run(section_data["heading"])
        heading_run.font.size = Pt(14)
        heading_run.bold = True
        
        # Body paragraphs
        for para_text in section_data["paragraphs"]:
            para = doc.add_paragraph(para_text)
            para.style = doc.styles["Normal"]
            para.paragraph_format.space_after = Pt(8)
            para.paragraph_format.space_before = Pt(4)
            for run in para.runs:
                run.font.size = Pt(12)
        
        doc.add_paragraph()  # section spacing
    
    # Save and verify
    doc.save(str(output_path))
    
    if not output_path.exists():
        return {"success": False, "error": "Document was not saved correctly"}
    
    file_size = output_path.stat().st_size
    if file_size == 0:
        return {"success": False, "error": "Document was saved but appears empty"}
    
    logger.info(f"Word document created: {output_path} ({file_size} bytes)")
    
    # Try to open with default application
    try:
        os.startfile(str(output_path))
        opened = True
    except Exception:
        opened = False
    
    return {
        "success": True,
        "message": f"Document '{filename}.docx' created successfully ({file_size // 1024}KB)",
        "path": str(output_path),
        "filename": f"{filename}.docx",
        "file_size_bytes": file_size,
        "opened_in_app": opened,
        "sections_created": len(content["sections"]),
    }


def _generate_document_content(topic: str, doc_type: str, custom_sections: Optional[list]) -> dict:
    """Generate structured document content."""
    
    if custom_sections:
        section_headings = custom_sections
    elif doc_type == "assignment":
        section_headings = ["Introduction", "Main Body", "Discussion", "Conclusion", "References"]
    elif doc_type == "report":
        section_headings = ["Executive Summary", "Introduction", "Findings", "Analysis", "Recommendations", "Conclusion"]
    elif doc_type == "essay":
        section_headings = ["Introduction", "Argument", "Counter-argument", "Analysis", "Conclusion"]
    else:
        section_headings = ["Introduction", "Overview", "Details", "Conclusion"]
    
    sections = []
    for heading in section_headings:
        sections.append({
            "heading": heading,
            "paragraphs": [_generate_section_text(heading, topic)],
        })
    
    return {
        "title": f"{topic}: A Comprehensive {doc_type.title()}",
        "sections": sections,
    }


def _generate_section_text(section: str, topic: str) -> str:
    """Generate placeholder text for each section. Phase 3 replaces this with Gemini."""
    templates = {
        "Introduction": f"This {section.lower()} provides an overview of {topic}. The following sections explore the key concepts, applications, and implications of this subject in detail.",
        "Main Body": f"{topic} encompasses a wide range of concepts and applications. This section examines the core principles and their practical significance.",
        "Discussion": f"The discussion of {topic} reveals important insights into its role and impact. Various perspectives help illuminate its broader significance.",
        "Conclusion": f"In conclusion, {topic} represents an important area of study. The key points covered in this document demonstrate its relevance and ongoing importance.",
        "References": "This section will contain academic references and citations. Please add appropriate sources from peer-reviewed journals and authoritative publications.",
        "Executive Summary": f"This document provides a comprehensive analysis of {topic}, covering its key aspects, implications, and recommendations for stakeholders.",
        "Findings": f"Research and analysis of {topic} reveals several significant findings that merit careful consideration.",
        "Analysis": f"A thorough analysis of {topic} demonstrates its complexity and the multifaceted nature of its applications and implications.",
        "Recommendations": f"Based on the analysis of {topic}, the following recommendations are proposed to address key challenges and opportunities.",
        "Overview": f"{topic} is a broad and important subject with numerous practical applications. This overview establishes the foundational concepts.",
        "Details": f"Detailed examination of {topic} reveals the intricate connections between its components and their real-world applications.",
    }
    return templates.get(section, f"This section covers the {section.lower()} aspects of {topic}. Further detailed content should be added based on specific requirements and research.")
