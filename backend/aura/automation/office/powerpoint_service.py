"""PowerPoint creation service using python-pptx."""
import os
import logging
from pathlib import Path
from typing import Optional
from datetime import datetime

logger = logging.getLogger("aura.office.powerpoint")


def get_output_dir() -> Path:
    output = Path.home() / "Documents" / "AURA"
    output.mkdir(parents=True, exist_ok=True)
    return output


async def create_presentation(
    topic: str,
    slide_count: int = 5,
    filename: Optional[str] = None,
) -> dict:
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt, Emu
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
    except ImportError:
        return {"success": False, "error": "python-pptx is not installed"}

    slide_count = max(2, min(slide_count, 30))
    
    if not filename:
        safe_topic = "".join(c if c.isalnum() or c in " _-" else "" for c in topic)[:40]
        filename = f"{safe_topic.strip()}_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    
    output_path = get_output_dir() / f"{filename}.pptx"
    
    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)
    
    # Generate slide outline
    slides_data = _generate_slide_outline(topic, slide_count)
    
    for idx, slide_data in enumerate(slides_data):
        if idx == 0:
            layout = prs.slide_layouts[0]  # Title slide
        else:
            layout = prs.slide_layouts[1]  # Title + Content
        
        slide = prs.slides.add_slide(layout)
        
        # Set background color
        from pptx.oxml.ns import qn
        from lxml import etree
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = RGBColor(0x0D, 0x1B, 0x2A)
        
        # Title
        if slide.shapes.title:
            slide.shapes.title.text = slide_data["title"]
            title_tf = slide.shapes.title.text_frame
            for para in title_tf.paragraphs:
                for run in para.runs:
                    run.font.color.rgb = RGBColor(0x95, 0x98, 0xFF)
                    run.font.bold = True
                    run.font.size = Pt(32 if idx == 0 else 28)
        
        # Content
        if len(slide.placeholders) > 1 and slide_data.get("content"):
            content_ph = slide.placeholders[1]
            tf = content_ph.text_frame
            tf.word_wrap = True
            
            # Clear and add bullets
            tf.text = ""
            first = True
            for bullet in slide_data["content"]:
                if first:
                    p = tf.paragraphs[0]
                    first = False
                else:
                    p = tf.add_paragraph()
                p.text = bullet
                p.level = 0
                for run in p.runs:
                    run.font.color.rgb = RGBColor(0xE0, 0xE4, 0xF0)
                    run.font.size = Pt(18)
        
        # Slide number (bottom right)
        txBox = slide.shapes.add_textbox(
            Inches(11.8), Inches(7.0), Inches(1.0), Inches(0.4)
        )
        tf = txBox.text_frame
        tf.text = f"{idx + 1}/{len(slides_data)}"
        for para in tf.paragraphs:
            for run in para.runs:
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0x60, 0x66, 0x80)
    
    prs.save(str(output_path))
    
    if not output_path.exists():
        return {"success": False, "error": "Presentation was not saved correctly"}
    
    file_size = output_path.stat().st_size
    actual_slides = len(prs.slides)
    
    logger.info(f"PowerPoint created: {output_path} ({actual_slides} slides, {file_size} bytes)")
    
    try:
        os.startfile(str(output_path))
        opened = True
    except Exception:
        opened = False
    
    return {
        "success": True,
        "message": f"Presentation '{filename}.pptx' created with {actual_slides} slides",
        "path": str(output_path),
        "filename": f"{filename}.pptx",
        "slide_count": actual_slides,
        "file_size_bytes": file_size,
        "opened_in_app": opened,
    }


def _generate_slide_outline(topic: str, count: int) -> list[dict]:
    """Generate slide content for the topic."""
    base_slides = [
        {"title": topic, "content": ["A comprehensive presentation"]},
        {"title": "Introduction", "content": [
            f"• Overview of {topic}",
            "• Key objectives of this presentation",
            "• Why this topic matters",
        ]},
        {"title": "Background", "content": [
            f"• Historical context of {topic}",
            "• Current state of the field",
            "• Fundamental concepts",
        ]},
        {"title": "Core Concepts", "content": [
            "• Key principles and definitions",
            "• Theoretical framework",
            "• Important terminology",
        ]},
        {"title": "How It Works", "content": [
            "• Step-by-step process",
            "• Key components and their roles",
            "• Technical overview",
        ]},
        {"title": "Applications", "content": [
            "• Real-world use cases",
            "• Industry applications",
            "• Practical examples",
        ]},
        {"title": "Advantages", "content": [
            "• Key benefits",
            "• Value proposition",
            "• Competitive advantages",
        ]},
        {"title": "Challenges", "content": [
            "• Current limitations",
            "• Open problems",
            "• Areas for improvement",
        ]},
        {"title": "Future Outlook", "content": [
            "• Emerging trends",
            "• Research directions",
            "• Expected developments",
        ]},
        {"title": "Conclusion", "content": [
            f"• Summary of {topic}",
            "• Key takeaways",
            "• Call to action",
        ]},
    ]
    
    if count <= len(base_slides):
        # Pick a balanced subset
        selected = [base_slides[0]]  # Always title
        step = (len(base_slides) - 1) / max(count - 1, 1)
        for i in range(1, count):
            idx = min(int(i * step), len(base_slides) - 1)
            selected.append(base_slides[idx])
        return selected[:count]
    else:
        return base_slides[:count]
